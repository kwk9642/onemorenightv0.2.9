#!/usr/bin/env python3
from pathlib import Path
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROWS = [
    # already in
    ("title.mp3（已有）", "Instrumental only, no vocals, no lyrics. A looping title theme for a text-based zombie shelter sim. Dusty American Midwest dusk, boarded farmhouse, analog radio hiss under a slow guitar and low cello. Melancholy, not epic, not heroic. Sparse piano after 30 seconds. Feels like one more night might be the last. 90-110 BPM, minor key, loopable 2 minutes."),
    ("hub.mp3（已有）", "Instrumental only, no vocals. Interior shelter ambience as music: wood creak, distant generator, muted folk guitar, analog tape warmth. Quiet colony management, Kentucky autumn, oil lamp, not a battle theme. Slow 70-85 BPM, looping, room tone mixed with music. Safe but never comfortable."),
    ("mission.mp3（已有）", "Instrumental only, no vocals. Stealth scavenger loop for walking a looted American town. Tension without drums at first, then a dry heartbeat kick. Empty mall air, fluorescent buzz as texture, detuned guitar, industrial pad. Not action Hollywood. Paranoia. 95 BPM, loopable."),
    ("combat.mp3（已有）", "Instrumental only, no vocals. Close-quarters zombie fight layer. Fast dry percussion, distorted bass, metallic hits, no orchestra swell, no choir. Ugly, physical, short loop 45-70 seconds so it can duck under the mission track. Relentless, not heroic."),
    ("occupy.mp3（已有）", "Instrumental only, no vocals. A held-breath theme for a stranger at the boarded door: stay, leave, or become an enemy. Slow pulse, radio voice almost forming words but never lyrics. Moral weight. 60 BPM, sparse piano and bowed bass."),
    ("end.mp3（已有）", "Instrumental only, no vocals. Permadeath end: the last survivor is gone. Not a movie credits fanfare. Empty house, wind in plywood, one guitar phrase that does not resolve. Cold. 2 minutes, fade on tape hiss."),
    # BGM still needed
    ("hub-night.mp3", "Instrumental only, no vocals. Same shelter as hub but after dark: everyone is asleep or afraid. Lower, slower, more wind against boards, distant dogs, one analog synth like a dying fridge. Loop for the overnight wait. 55-65 BPM."),
    ("hub-hope.mp3", "Instrumental only, no vocals. Rare good morning in the wasteland: food for a week, nobody bit. Warm analog keys, soft acoustic guitar, still dusty, not cheerful pop. A crack of light through nailed plywood. Loopable, 80 BPM."),
    ("hub-starving.mp3", "Instrumental only, no vocals. The cupboards are empty. Stomach-cramp drone, dry mouth percussion of clock ticks, no melody for a long time, then a thin harmonica far away. Anxiety loop for a hungry colony. No vocals."),
    ("hub-sick.mp3", "Instrumental only, no vocals. Someone in the house is bitten or fevered. Medical horror without jumpscares: slow dripping, detuned piano, analog tape wobble, heartbeat very quiet. Loop. Do not use choir."),
    ("mission-house.mp3", "Instrumental only, no vocals. Looting a suburban American house. Carpet muffled steps as rhythm, jewelry-box music box almost broken, dust motes, not a jump-scare track. Intimate, claustrophobic, loopable."),
    ("mission-mall.mp3", "Instrumental only, no vocals. Dead shopping mall. Huge reverb, distant cart rattling, dead Muzak ghost as a texture not a tune, escalator hum. Horror of retail silence. Loop 2 min."),
    ("mission-hospital.mp3", "Instrumental only, no vocals. Abandoned hospital. Fluorescent stutter as percussion, rubber-shoe squeak, very high sine that never becomes a melody. Clinical dread. No vocals, no scream samples that sound comic."),
    ("mission-warehouse.mp3", "Instrumental only, no vocals. Night warehouse. Metal racks, forklift ghost, rain on corrugated roof, industrial bass. Workers long gone. Loopable stealth."),
    ("mission-pd.mp3", "Instrumental only, no vocals. Looting a police station. Cold concrete, buzzing evidence cage, distant alarm that died years ago still in the key. Tension, authority rotting. No heroic cop theme."),
    ("mission-farm.mp3", "Instrumental only, no vocals. Rural Kentucky farm at dusk. Crow, dry corn, barn wood, banjo used as texture not folk song, wind through silo. Isolation. Loop."),
    ("rain.mp3", "Instrumental ambient loop, no vocals, no melody. Heavy Midwest rain on a boarded roof, gutter drip, thunder very far. Pure weather bed to layer under hub or mission. 2 minutes seamless."),
    ("wind-night.mp3", "Instrumental ambient loop, no vocals. Night wind hitting plywood and plastic sheeting, house flexing, one loose nail ticking. No music. Sound-design bed."),
    ("radio-static.mp3", "Instrumental / sound design, no sung lyrics. Shortwave radio bed: hiss, numbers-station pulses, a voice that is never intelligible, analog knobs. Loop under occupy or title. Unsettling but quiet."),
    # stingers
    ("sting-morning.mp3", "Instrumental stinger 8-12 seconds, no vocals. Dawn through cracks in boards. One guitar harmonic, a bird too far, then silence. Used when overnight succeeds. Not a fanfare."),
    ("sting-wipe.mp3", "Instrumental stinger 10-15 seconds, no vocals. The last person in the house is dead. Low brass cluster, tape stop, wood slam. Permadeath. No choir, no scream."),
    ("sting-recruit.mp3", "Instrumental stinger 8 seconds, no vocals. A living stranger is allowed in. Cautious warmth, door bar lifted, one major interval that immediately clouds. Not victory music."),
    ("sting-turnaway.mp3", "Instrumental stinger 8 seconds, no vocals. You sent them back into the dark. Guilt. Door bars drop, footsteps recede into gravel. Sparse."),
    ("sting-enemy.mp3", "Instrumental stinger 8 seconds, no vocals. The person at the door draws a weapon. Fast analog stab, metal, then cut to silence."),
    ("sting-rare.mp3", "Instrumental stinger 6-8 seconds, no vocals. You found something rare in the rubble: a plate carrier, a rifle, medicine. Quiet awe, not slot-machine jingle. Analog chime through dust."),
    ("sting-legendary.mp3", "Instrumental stinger 10 seconds, no vocals. Legendary find. Still not a game-show sting. Slow brass and guitar harmonic, tape saturation, weight. One More Night tone."),
    ("sting-infected.mp3", "Instrumental stinger 8 seconds, no vocals. Someone came home bitten. Dread realization. Heartbeat, then a sour piano note that does not resolve."),
    ("sting-lowfood.mp3", "Instrumental stinger 6 seconds, no vocals. Pantry is a day from empty. Stomach hollow tone, cupboard click. Warning, not comic."),
    ("sting-walls.mp3", "Instrumental stinger 8 seconds, no vocals. Something hit the outer wall at night. Wood boom, then listening. Horror of the house as a drum."),
    ("sting-extract.mp3", "Instrumental stinger 8-12 seconds, no vocals. The squad made it home. Exhausted relief, gate chain, bar dropped. Not triumph."),
    ("sting-newrun.mp3", "Instrumental stinger 10 seconds, no vocals. New survivor, new house, same world. A reset that is not hopeful. Paper page, match strike, one guitar."),
    ("sting-shop.mp3", "Instrumental stinger 4-6 seconds, no vocals. Opening the paper-and-frame market. Dry paper rustle turned into a tiny motif. No cash-register sound."),
    # UI / foley-like musical SFX (Suno as short design)
    ("ui-click.mp3", "Ultra short instrumental hit 1-2 seconds, no vocals. Soft wooden switch, analog button, paper tap. UI click for a journal game. Not a plastic mobile click. Export and we will trim."),
    ("ui-confirm.mp3", "Short instrumental 2 seconds, no vocals. Low guitar pluck plus a tiny analog beep, confirming a choice in a dark shelter sim. Soft, dry, no sparkle."),
    ("ui-back.mp3", "Short instrumental 1-2 seconds, no vocals. Page turning back in a field notebook. Paper, not whoosh."),
    ("ui-error.mp3", "Short instrumental 2 seconds, no vocals. Cannot go out, too broken, too sad. Muted warning buzzer through wood. Not arcade error."),
    ("ui-page.mp3", "Short foley-music 2 seconds, no vocals. Thick paper page in a damp notebook. For log entries."),
    ("door-bar.mp3", "Short sound-design 3 seconds, no vocals. Heavy timber bar dropped across a farmhouse door. Wood, iron bracket, finality."),
    ("door-open.mp3", "Short sound-design 3 seconds, no vocals. Boarded interior door opening on dry hinges. Caution. No cartoon creak."),
    ("door-knock.mp3", "Short sound-design 4 seconds, no vocals. Three knocks on plywood from outside at night. Human, tired, maybe not human. Silence after."),
    ("foot-wood.mp3", "Loopable 8-second footsteps on old farmhouse boards, no music, no vocals. Slow, careful, weight shifting. We will slice."),
    ("foot-gravel.mp3", "Loopable 8-second footsteps on Midwest gravel and glass, no vocals. Scavenger pace."),
    ("foot-mall.mp3", "Loopable 8-second sneakers on dead mall tile, huge quiet reverb, no vocals."),
    ("search-drawer.mp3", "Short 3 seconds, no vocals. Rummaging a kitchen drawer in the dark. Cutlery, paper, a useful clink at the end."),
    ("search-crate.mp3", "Short 3 seconds, no vocals. Crowbar on a wooden crate, lid, rustle of packing. Loot start."),
    ("pickup-common.mp3", "Short 2 seconds, no vocals. Picking up a can, a knife, a shirt. Soft, physical, not a coin sound."),
    ("pickup-meds.mp3", "Short 2 seconds, no vocals. Plastic pill bottle, foil, medical weight. Small hope."),
    ("pickup-ammo.mp3", "Short 2 seconds, no vocals. A handful of cartridges in a palm. Metal, not a reload animation."),
    ("swing-melee.mp3", "Short 1 second, no vocals. Crowbar or bat swing through air. Whoosh of cheap wood and panic. Not anime."),
    ("impact-zombie.mp3", "Short 1 second, no vocals. Blunt tool hitting something that used to be a person. Ugly, dry, not splat-comedy."),
    ("impact-wood.mp3", "Short 1 second, no vocals. Weapon hitting a doorframe in the dark. Missed the dead."),
    ("gun-pistol.mp3", "Short 1-2 seconds, no vocals. One indoor pistol shot, Midwestern house, ringing after. Brutal, not action-movie crispy."),
    ("gun-shotgun.mp3", "Short 2 seconds, no vocals. Pump shotgun in a hallway. Room boom, ears ringing. One shot only."),
    ("gun-rifle.mp3", "Short 2 seconds, no vocals. Hunting rifle outdoors at dusk, echo off a barn. One shot."),
    ("hurt-player.mp3", "Short 2 seconds, no vocals. Player or survivor is hit. Breath knock, not a Wilhelm scream. Human."),
    ("death-down.mp3", "Short 3 seconds, no vocals. A body going down on wood. Final. No choir."),
    ("zombie-idle.mp3", "Loop 10-15 seconds, no music, no vocals. Distant undead: dry throat, shuffle, not a movie horde roar. One or two bodies outside the wall."),
    ("zombie-notice.mp3", "Short 3 seconds, no vocals. They heard you. Attention snap, wet breath closer."),
    ("zombie-attack.mp3", "Short 2 seconds, no vocals. Bite or grab. Fast, ugly, close mic."),
    ("zombie-down.mp3", "Short 2 seconds, no vocals. Dead thing stops. Weight, not a point-score."),
    ("glass-break.mp3", "Short 2 seconds, no vocals. A pane in a boarded window finally goes. Night air rushes in."),
    ("generator.mp3", "Loop 20 seconds, no vocals. Dying gasoline generator outside the hardware store. Uneven idle. Layer under hub."),
    ("crow.mp3", "Short 3 seconds, no vocals. One American crow over a dead field. Isolation."),
    ("match.mp3", "Short 2 seconds, no vocals. Match strike in a dark kitchen. Tiny fire. Title or overnight."),
    ("radio-voice.mp3", "8-12 seconds, no sung lyrics. A broken radio almost says a name, then static. For recruit-from-radio events. Spoken fragments unintelligible on purpose."),
    ("build-hammer.mp3", "Short 3 seconds, no vocals. Hammering a board over a window from inside. Shelter upgrade."),
    ("saw-wood.mp3", "Short 4 seconds, no vocals. Hand saw on lumber in a quiet house. Upgrade / craft."),
    ("cook.mp3", "Short 4 seconds, no vocals. Pan on a scavenged stove, oil, a lid. Domestic and fragile."),
    ("water-pour.mp3", "Short 3 seconds, no vocals. Pouring precious water into a metal cup. Thirst."),
    ("sleep.mp3", "8 seconds, no vocals. House settling as people try to sleep. Bed springs, one cough, then wind. Overnight start."),
    ("clock-tick.mp3", "Loop 10 seconds, no vocals. A cheap wall clock in a dead hardware store. Time as threat."),
    ("chain-gate.mp3", "Short 3 seconds, no vocals. Unlocking a chain on a farm gate before a run. Going out."),
    ("backpack.mp3", "Short 2 seconds, no vocals. Shoulder a heavy pack. Straps, cans, breath."),
    ("bandage.mp3", "Short 3 seconds, no vocals. Tearing cloth, wrapping a wound. Medical, intimate."),
    ("write-log.mp3", "Short 3 seconds, no vocals. Pencil on damp paper. The day's log. Journal identity."),
    ("lighter.mp3", "Short 1 second, no vocals. Cheap metal lighter, failed then lit. Nervous hands."),
    ("can-open.mp3", "Short 2 seconds, no vocals. Manual can opener on a rusted tin. Food."),
    ("nail.mp3", "Short 2 seconds, no vocals. One nail driven into plywood. Fortify."),
    ("breath-cold.mp3", "Loop 8 seconds, no vocals. Close-mic breathing in a cold hallway. Stealth layer."),
    ("crowd-far.mp3", "Loop 15 seconds, no vocals. A horde very far away, like weather. Never close. Threat on the map."),
    ("train-ghost.mp3", "10 seconds, no vocals. A freight line that does not run anymore: metal memory, one horn almost. Midwest haunt. Stinger or occupy."),
    ("church-bell-dead.mp3", "8 seconds, no vocals. One cracked bell very far. Sunday in a town of the dead. Not holy, just empty."),
    ("insect-night.mp3", "Loop 20 seconds, no vocals. Kentucky night insects, then they stop. The stop is the horror. Layer under hub-night."),
    ("flare.mp3", "Short 4 seconds, no vocals. Highway flare struck on wet asphalt. Red light, chemical hiss. Rescue that is not coming."),
    ("tape.mp3", "Short 2 seconds, no vocals. Duct tape pulled and torn. Repair, wounds, windows."),
    ("keys.mp3", "Short 2 seconds, no vocals. A ring of hardware-store keys. Which door still opens."),
    ("safe-click.mp3", "Short 3 seconds, no vocals. A cheap gun safe finally opening. Tension release, not treasure jingle."),
    ("elevator-dead.mp3", "Short 4 seconds, no vocals. High-rise elevator that will never move. Cable tick in the shaft. Vertical dread."),
    ("title-sting.mp3", "6-8 second instrumental logo sting, no vocals. For the words One More Night. Guitar and tape, not an orchestra hit. Same world as title.mp3."),
]

NOTE = (
    "Suno 設定建議：Custom / Instrumental。Style 欄可再加：dark folk, analog horror, "
    "found-footage, Midwest autumn, no vocals, no lyrics, lo-fi tape。\n"
    "出糧後請匯出 MP3，檔名改成左欄（去掉「已有」）。Loop 的曲子留 1–2 分鐘；短的 stinger／音效請自己在 2–8 秒處剪。\n"
    "Suno 較擅長曲子與 stinger，純腳步／槍聲可能要多 roll 幾次，或當「聲音設計短句」來寫。"
)


def set_run(run, size=11, bold=False, name="Noto Serif CJK TC"):
    run.font.size = Pt(size)
    run.bold = bold
    run.font.name = name
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), name)
    rFonts.set(qn("w:ascii"), "Calibri")
    rFonts.set(qn("w:hAnsi"), "Calibri")


def add_p(doc, text, size=11, bold=False, space_after=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    run = p.add_run(text)
    set_run(run, size=size, bold=bold)
    return p


def main():
    have = sum(1 for n, _ in ROWS if "已有" in n)
    need = len(ROWS) - have
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(1.6)
    sec.right_margin = Cm(1.6)
    sec.top_margin = Cm(1.6)
    sec.bottom_margin = Cm(1.6)

    add_p(doc, "再活一晚 One More Night", size=20, bold=True, space_after=2)
    add_p(doc, "v0.1　Suno 音效／配樂清單", size=14, bold=True, space_after=10)
    add_p(
        doc,
        f"現成可播：{have} 首 BGM（title / hub / mission / combat / occupy / end）。"
        f"建議再補：{need} 個檔。合計 {len(ROWS)} 列。缺的越多越好，慢慢補，檔名對上就能接。",
        size=11,
        space_after=8,
    )
    add_p(doc, NOTE, size=10, space_after=12)

    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(5.2)
    table.columns[1].width = Cm(12.6)

    hdr = table.rows[0].cells
    hdr[0].text = ""
    hdr[1].text = ""
    for cell, text in ((hdr[0], "檔名（左）"), (hdr[1], "Suno 指令，整段複製（右）")):
        cell.paragraphs[0].clear()
        run = cell.paragraphs[0].add_run(text)
        set_run(run, size=11, bold=True)
        from docx.oxml import OxmlElement

        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "161410")
        shd.set(qn("w:val"), "clear")
        cell._tc.get_or_add_tcPr().append(shd)
        run.font.color.rgb = RGBColor(0xE8, 0xE0, 0xD4)

    for name, prompt in ROWS:
        row = table.add_row().cells
        row[0].paragraphs[0].clear()
        r0 = row[0].paragraphs[0].add_run(name)
        set_run(r0, size=10, bold=True)
        row[1].paragraphs[0].clear()
        r1 = row[1].paragraphs[0].add_run(prompt)
        set_run(r1, size=9, bold=False)
        for p in row[0].paragraphs + row[1].paragraphs:
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.space_before = Pt(4)

    add_p(doc, "", size=11, space_after=6)
    add_p(
        doc,
        "已有六首不必重做；若你更滿意新 roll，用同檔名覆蓋 public/game/music/ 即可。短音效日後會進 public/game/sfx/。交給我時請用這個檔名。",
        size=10,
        space_after=4,
    )

    out_dir = Path("/workspace/public")
    out_dir.mkdir(exist_ok=True)
    path = out_dir / "One-More-Night-Suno-音效清單.docx"
    doc.save(path)
    Path("/workspace/artifacts").mkdir(exist_ok=True)
    doc.save(Path("/workspace/artifacts") / path.name)
    print(path, path.stat().st_size, "rows", len(ROWS), "need", need)


if __name__ == "__main__":
    main()
